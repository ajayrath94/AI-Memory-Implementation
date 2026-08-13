"""
EXPORT — download memory data (Track 1 + Track 2) as an Excel workbook, for
continuous-improvement inspection. On-demand, always current.

GET /export/memory-xlsx            all users
GET /export/memory-xlsx?user_id=X  one user
"""
import io, json
from fastapi import APIRouter
from fastapi.responses import StreamingResponse

router = APIRouter()

_HEADER_FILL = "2C3E50"
_FONT = "Arial"


def _style_header(ws, ncols):
    from openpyxl.styles import Font, PatternFill, Alignment
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name=_FONT, bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _autowidth(ws, widths):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


@router.get("/memory-xlsx")
def export_memory_xlsx(user_id: str = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from supabase_store import get_client
    db = get_client()

    def q(table, select="*"):
        query = db.table(table).select(select)
        if user_id:
            query = query.eq("user_id", user_id)
        try:
            return query.execute().data or []
        except Exception as e:
            print(f"[Export] {table} query failed: {e}")
            return []

    wb = Workbook()

    # ---- Sheet 1: Track 1 narrative ----
    ws1 = wb.active
    ws1.title = "Track1_Narrative"
    h1 = ["user_id", "session_count", "summary", "key_facts", "updated_at"]
    ws1.append(h1)
    for um in q("user_memory", "user_id,session_count,summary,key_facts,updated_at"):
        kf = um.get("key_facts")
        if isinstance(kf, list):
            kf = " | ".join(str(x) for x in kf)
        ws1.append([
            um.get("user_id"), um.get("session_count"),
            um.get("summary"), kf, um.get("updated_at"),
        ])
    _style_header(ws1, len(h1))
    _autowidth(ws1, [20, 12, 80, 60, 20])
    for row in ws1.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=_FONT, size=10)
            c.alignment = c.alignment.copy(wrap_text=True, vertical="top")

    # ---- Sheet 2: Track 2 clusters ----
    ws2 = wb.create_sheet("Track2_Clusters")
    h2 = ["user_id", "label", "pillar", "strength", "event_count", "trend", "status", "attributes"]
    ws2.append(h2)
    clusters = q("interest_clusters",
                 "user_id,label,pillar,strength,event_count,trend,status,attributes")
    clusters.sort(key=lambda c: (str(c.get("user_id")), -(c.get("strength") or 0)))
    for cl in clusters:
        attrs = cl.get("attributes")
        if isinstance(attrs, dict):
            attrs = json.dumps(attrs, ensure_ascii=False)
        ws2.append([
            cl.get("user_id"), cl.get("label"), cl.get("pillar"),
            round(cl.get("strength") or 0, 3), cl.get("event_count"),
            cl.get("trend"), cl.get("status"), attrs,
        ])
    _style_header(ws2, len(h2))
    _autowidth(ws2, [20, 26, 16, 10, 12, 12, 10, 40])
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=_FONT, size=10)

    # ---- Sheet 3: Sessions ----
    ws3 = wb.create_sheet("Sessions")
    h3 = ["user_id", "session_id", "created_at", "summary"]
    ws3.append(h3)
    sess = q("sessions", "user_id,id,created_at,summary")
    sess.sort(key=lambda s: (str(s.get("user_id")), str(s.get("created_at"))))
    for s in sess:
        ws3.append([s.get("user_id"), s.get("id"), s.get("created_at"), s.get("summary")])
    _style_header(ws3, len(h3))
    _autowidth(ws3, [20, 38, 20, 80])
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=_FONT, size=10)
            c.alignment = c.alignment.copy(wrap_text=True, vertical="top")

    # ---- Sheet 4: Messages (the diagnostic layer under both tiers) ----
    # Text + classification, NOT the embedding/pillar_vector (unreadable arrays).
    # messages are keyed by session_id, so map session -> user_id first.
    ws4 = wb.create_sheet("Messages")
    h4 = ["user_id", "role", "content", "model", "timestamp",
          "core", "emotion", "functional", "modifiers", "score",
          "memory_tier", "is_summary"]
    ws4.append(h4)
    # build session_id -> user_id map (scoped to the export's user filter)
    sess_rows = q("sessions", "id,user_id")
    sess_map = {s["id"]: s.get("user_id") for s in sess_rows}
    # pull messages for those sessions
    msg_rows = []
    try:
        if user_id:
            sids = [sid for sid, uid in sess_map.items() if uid == user_id]
        else:
            sids = list(sess_map.keys())
        # chunk the IN filter to avoid overly long URLs
        for i in range(0, len(sids), 50):
            chunk = sids[i:i+50]
            if not chunk:
                continue
            r = (db.table("messages").select(
                    "session_id,role,content,model,timestamp,pillar_core,"
                    "pillar_emotion,pillar_functional,pillar_modifiers,"
                    "pillar_score,memory_tier,is_summary")
                 .in_("session_id", chunk).execute()).data or []
            msg_rows.extend(r)
    except Exception as e:
        print(f"[Export] messages query failed: {e}")
    # sort by user then time
    msg_rows.sort(key=lambda m: (str(sess_map.get(m.get("session_id"))),
                                 str(m.get("timestamp"))))
    for m in msg_rows:
        mods = m.get("pillar_modifiers")
        if isinstance(mods, list):
            mods = ", ".join(str(x) for x in mods)
        ws4.append([
            sess_map.get(m.get("session_id")), m.get("role"),
            m.get("content"), m.get("model"), m.get("timestamp"),
            m.get("pillar_core"), m.get("pillar_emotion"),
            m.get("pillar_functional"), mods,
            round(m.get("pillar_score") or 0, 3),
            m.get("memory_tier"), m.get("is_summary"),
        ])
    _style_header(ws4, len(h4))
    _autowidth(ws4, [20, 10, 70, 16, 20, 14, 12, 14, 20, 8, 12, 10])
    from openpyxl.styles import Font as _F
    for row in ws4.iter_rows(min_row=2):
        for cell in row:
            cell.font = _F(name=_FONT, size=10)
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"nancy_memory_{user_id or 'all'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
